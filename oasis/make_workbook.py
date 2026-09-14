"""Generate the synthetic maintenance workbook. Run: python make_workbook.py

A spreadsheet, not a document. It goes into SQLite and is queried exactly -
never chunked into a vector store, because chunking separates headers from rows
and loses the structure that makes a table worth having.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = "corpus"
os.makedirs(OUT, exist_ok=True)

WORK_ORDERS = [
    ("WO-24118", "2024-04-12", "P-101", "Mechanical seal replaced, routine",
     "Preventive", 14, 184000, "K. Nair", "Closed"),
    ("WO-24237", "2024-07-03", "P-102", "Bearing greased, routine",
     "Preventive", 2, 4500, "S. Rao", "Closed"),
    ("WO-24406", "2024-10-16", "P-101", "Mechanical seal replaced, leakage",
     "Breakdown", 18, 184000, "K. Nair", "Closed"),
    ("WO-24512", "2024-12-08", "V-201", "Relief valve tested",
     "Statutory", 6, 32000, "A. Menon", "Closed"),
    ("WO-25077", "2025-03-02", "P-101", "Bearing replaced, high vibration",
     "Breakdown", 22, 26500, "S. Rao", "Closed"),
    ("WO-25130", "2025-04-18", "P-102", "Impeller inspected",
     "Preventive", 8, 12000, "K. Nair", "Closed"),
    ("WO-25214", "2025-06-19", "P-101", "Mechanical seal replaced, leakage",
     "Breakdown", 16, 184000, "K. Nair", "Closed"),
    ("WO-25301", "2025-08-05", "HE-45B", "Tube bundle cleaned",
     "Preventive", 36, 210000, "A. Menon", "Closed"),
    ("WO-25390", "2025-09-21", "P-101", "Seal flush strainer cleaned",
     "Corrective", 4, 3200, "S. Rao", "Closed"),
    ("WO-25447", "2025-11-02", "V-201", "Gasket replaced",
     "Corrective", 9, 18000, "A. Menon", "Closed"),
    ("WO-25502", "2025-12-24", "P-101", "Mechanical seal replaced, leakage",
     "Breakdown", 19, 184000, "K. Nair", "Closed"),
    ("WO-26041", "2026-03-28", "P-101", "Mechanical seal replaced, leakage",
     "Breakdown", 17, 184000, "K. Nair", "Closed"),
    ("WO-26088", "2026-05-14", "P-102", "Coupling replaced",
     "Corrective", 11, 48000, "S. Rao", "Closed"),
    ("WO-26129", "2026-06-30", "P-101", "Coupling alignment check",
     "Preventive", 6, 9000, "S. Rao", "Closed"),
    ("WO-26174", "2026-08-11", "HE-45B", "Leak repair on shell side",
     "Breakdown", 28, 96000, "A. Menon", "Open"),
]

HEADERS = ["Work Order", "Date", "Equipment", "Description", "Type",
           "Downtime Hours", "Cost Rs", "Engineer", "Status"]

wb = Workbook()
ws = wb.active
ws.title = "Work Orders"
ws.append(HEADERS)

head_fill = PatternFill("solid", fgColor="E9E2F3")
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = head_fill

for row in WORK_ORDERS:
    ws.append(list(row))

for col in ws.columns:
    width = max(len(str(c.value or "")) for c in col)
    ws.column_dimensions[col[0].column_letter].width = min(width + 3, 40)

ws.freeze_panes = "A2"

path = os.path.join(OUT, "maintenance_workbook.xlsx")
wb.save(path)
print(f"written {path}  ({len(WORK_ORDERS)} rows)")
