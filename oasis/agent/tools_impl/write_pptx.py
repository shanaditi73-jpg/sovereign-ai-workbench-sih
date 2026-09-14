"""Tool: produce a PowerPoint deck. STATE-CHANGING - needs approval.

Title slide, one slide per section, and a closing sources slide naming every
document and page the content came from.
"""

import datetime

from agent.tools_impl._common import OUTPUT_DIR




def _normalise(slides):
    """Small models produce slides in several shapes. Accept them all.

    [{"heading": "...", "bullets": [...]}]   the intended form
    [{"title": "...", "points": [...]}]      near misses on key names
    [{"heading": "...", "bullets": "a\nb"}]  bullets as one string
    ["Heading: a; b", ...]                   plain strings
    """
    out = []
    for item in (slides or []):
        if isinstance(item, str):
            head, _, rest = item.partition(":")
            bullets = [b.strip() for b in rest.replace("\n", ";").split(";")
                       if b.strip()]
            out.append({"heading": head.strip() or "Slide",
                        "bullets": bullets})
            continue
        if not isinstance(item, dict):
            continue
        head = (item.get("heading") or item.get("title")
                or item.get("header") or "Slide")
        bl = (item.get("bullets") or item.get("points")
              or item.get("content") or item.get("body") or [])
        if isinstance(bl, str):
            bl = [b.strip() for b in bl.replace(";", "\n").split("\n")
                  if b.strip()]
        out.append({"heading": str(head),
                    "bullets": [str(b) for b in bl if str(b).strip()]})
    return out


def tool_write_pptx(title: str, subtitle: str = "", slides: list = None,
                    sources: list = None, filename: str = None,
                    author: str = "", **_ignored):
    """Produce a presentation. STATE-CHANGING - needs approval first.

    slides is a list of {"heading": str, "bullets": [str, ...]}
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    INK = RGBColor(0x1A, 0x16, 0x30)
    MUTED = RGBColor(0x5F, 0x57, 0x80)
    ACCENT = RGBColor(0x6A, 0x57, 0xC4)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def textbox(slide, x, y, w, h, text, size, bold=False, colour=INK):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.runs[0].font.size = Pt(size)
        p.runs[0].font.bold = bold
        p.runs[0].font.color.rgb = colour
        return tf

    # --- title slide
    s0 = prs.slides.add_slide(blank)
    textbox(s0, 0.9, 2.5, 11.5, 1.4, title or "Report", 40, True)
    if subtitle:
        textbox(s0, 0.9, 3.9, 11.5, 0.8, subtitle, 17, False, MUTED)
    footer = "Prepared on premises from plant documentation"
    if author:
        footer += f"  ·  {author}"
    textbox(s0, 0.9, 6.5, 11.5, 0.5, footer, 11, False, MUTED)

    # --- content slides
    for item in _normalise(slides):
        sl = prs.slides.add_slide(blank)
        textbox(sl, 0.9, 0.7, 11.5, 0.9, item.get("heading", ""), 28, True)

        tb = sl.shapes.add_textbox(Inches(0.9), Inches(1.9),
                                   Inches(11.5), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for b in item.get("bullets", []):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = str(b)
            p.space_after = Pt(14)
            p.runs[0].font.size = Pt(16)
            p.runs[0].font.color.rgb = INK

    if len(prs.slides._sldIdLst) == 1 and subtitle:
        sl = prs.slides.add_slide(blank)
        textbox(sl, 0.9, 0.7, 11.5, 0.9, "Summary", 28, True)
        tb = sl.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(4))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text = subtitle
        p.runs[0].font.size = Pt(16)

    # --- sources slide
    if sources:
        sl = prs.slides.add_slide(blank)
        textbox(sl, 0.9, 0.7, 11.5, 0.9, "Sources", 28, True, ACCENT)
        tb = sl.shapes.add_textbox(Inches(0.9), Inches(1.9),
                                   Inches(11.5), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, src in enumerate(sources, 1):
            pg = f", page {src['page']}" if src.get("page") is not None else ""
            p = tf.paragraphs[0] if i == 1 else tf.add_paragraph()
            p.text = f"[{i}]  {src['document']}{pg}"
            p.space_after = Pt(10)
            p.runs[0].font.size = Pt(13)
            p.runs[0].font.color.rgb = MUTED

    name = filename or f"deck_{datetime.datetime.now():%H%M%S}.pptx"
    path = OUTPUT_DIR / name
    prs.save(path)
    return {"file": name, "path": str(path), "slides": len(prs.slides._sldIdLst)}
